'''Offline unit tests for LedgerLens (no API keys, no internet, no embedding model).

Run from the project folder:   python -m unittest discover -s tests
'''
from __future__ import annotations

import io
import json
import logging
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from langchain_core.documents import Document  # noqa: E402
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel  # noqa: E402
from langchain_core.messages import AIMessage  # noqa: E402
from pypdf import PdfWriter  # noqa: E402

import agent as agent_module  # noqa: E402
import ingestion  # noqa: E402
import llm  # noqa: E402
import retriever  # noqa: E402
from agent import (AgentAnswer, AgentFailure, AnswerClaim, DraftReply, EvidenceItem,  # noqa: E402
                   LedgerAgent, NO_ANSWER_TEXT, PassageLedger, REFUSALS, build_search_tool)
from ingestion import DocumentRejectedError, read_document  # noqa: E402
from llm import Provider  # noqa: E402
from retriever import SearchHit  # noqa: E402

SAMPLES = ROOT / 'sample_documents'
BAD_FILES = SAMPLES / 'error_test_files'


def setUpModule() -> None:
    logging.disable(logging.CRITICAL)


def tearDownModule() -> None:
    logging.disable(logging.NOTSET)


# ----------------------------------------------------------------------------- file builders
def blank_pdf() -> bytes:
    '''A valid PDF whose single page has no text at all (what a scanned page looks like to a parser).'''
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def locked_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('open-sesame')
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


CLAIMS_CSV = (
    'number,department,amount,state\n'
    'R-1,Finance,41825,Submitted\n'
    'R-2,HR,19938,Declined\n'
    'R-3,IT,57593,Pending\n'
).encode('utf-8')


class RejectedTypeTests(unittest.TestCase):
    '''Files of a type we never read must be refused with a specific, helpful explanation.'''

    def assertRejected(self, name: str, data: bytes, *fragments: str) -> str:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document(name, data)
        message = str(caught.exception)
        for fragment in fragments:
            self.assertIn(fragment, message)
        return message

    def test_word_document_points_to_save_as(self) -> None:
        self.assertRejected('report.docx', b'PK\x03\x04', 'Save As', 'PDF')

    def test_zip_archive_says_unpack_first(self) -> None:
        self.assertRejected('bundle.zip', b'PK\x03\x04', 'Unpack')

    def test_old_excel_asks_for_xlsx_or_csv(self) -> None:
        self.assertRejected('budget.xls', b'\xd0\xcf\x11\xe0', '.xls', '.xlsx')

    def test_picture_explains_missing_ocr(self) -> None:
        self.assertRejected('receipt.png', b'\x89PNG', 'OCR')

    def test_slides_point_to_pdf_export(self) -> None:
        self.assertRejected('deck.pptx', b'PK', 'PDF')

    def test_unknown_extension_is_named(self) -> None:
        self.assertRejected('notes.xyz', b'abc', "'.xyz'")

    def test_missing_extension_is_explained(self) -> None:
        self.assertRejected('README', b'abc', 'no extension')

    def test_type_is_checked_before_emptiness(self) -> None:
        self.assertRejected('empty.docx', b'', 'cannot be used')

    def test_every_message_lists_the_allowed_types(self) -> None:
        for name in ('a.docx', 'a.zip', 'a.xls', 'a.png', 'a.pptx', 'a.xyz', 'a'):
            self.assertIn(ingestion.ALLOWED_TYPES_NOTE, self.assertRejected(name, b'x'))


class BrokenFileTests(unittest.TestCase):
    '''Empty, blank, corrupt, locked or disguised files: skipped or explained, never a crash.'''

    def test_zero_byte_file_gives_nothing(self) -> None:
        self.assertEqual(read_document('empty.txt', b''), [])

    def test_whitespace_only_text_gives_nothing(self) -> None:
        self.assertEqual(read_document('blank.txt', b'  \n\t \n'), [])

    def test_header_only_csv_gives_nothing(self) -> None:
        self.assertEqual(read_document('headers.csv', b'number,department,amount\n'), [])

    def test_pdf_without_text_gives_nothing(self) -> None:
        self.assertEqual(read_document('scan.pdf', blank_pdf()), [])

    def test_corrupt_pdf_gets_the_pdf_message(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('broken.pdf', b'%PDF-1.4\nnot really a pdf\n')
        self.assertIn('PDF could not be opened', str(caught.exception))

    def test_password_protected_pdf_is_named_as_such(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('locked.pdf', locked_pdf())
        self.assertIn('password', str(caught.exception))

    def test_text_renamed_to_xlsx_gets_the_excel_message(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('fake.xlsx', b'just plain text\n')
        self.assertIn('Excel file could not be opened', str(caught.exception))

    def test_zip_renamed_to_xlsx_is_refused(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('hello.txt', 'hi')
        with self.assertRaises(DocumentRejectedError):
            read_document('hidden.xlsx', buffer.getvalue())

    def test_binary_data_named_txt_is_refused(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('data.txt', b'abc\x00\x01\x02def')
        self.assertIn('text file could not be read', str(caught.exception))

    def test_binary_data_named_csv_is_refused(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('data.csv', b'a,b\n\x00\x00,1\n')
        self.assertIn('CSV file could not be read', str(caught.exception))

    def test_oversize_file_is_refused(self) -> None:
        with mock.patch.object(ingestion, 'MAX_FILE_BYTES', 10):
            with self.assertRaises(DocumentRejectedError) as caught:
                read_document('big.txt', b'0123456789ABC')
        self.assertIn('20 MiB', str(caught.exception))

    def test_too_many_characters_is_refused(self) -> None:
        with mock.patch.object(ingestion, 'MAX_EXTRACTED_CHARACTERS', 20):
            with self.assertRaises(DocumentRejectedError) as caught:
                read_document('long.txt', b'x' * 50)
        self.assertIn('over the limit', str(caught.exception))

    def test_too_many_csv_records_is_refused(self) -> None:
        with mock.patch.object(ingestion, 'MAX_RECORDS', 2):
            with self.assertRaises(DocumentRejectedError):
                read_document('rows.csv', CLAIMS_CSV)

    def test_too_many_excel_rows_is_refused(self) -> None:
        data = workbook_bytes({'S': [['a', 'b']] + [[i, i] for i in range(6)]})
        with mock.patch.object(ingestion, 'MAX_RECORDS', 3):
            with self.assertRaises(DocumentRejectedError) as caught:
                read_document('rows.xlsx', data)
        self.assertIn('20,000 rows', str(caught.exception))

    def test_error_messages_hide_internal_details(self) -> None:
        with self.assertRaises(DocumentRejectedError) as caught:
            read_document('broken.pdf', b'%PDF-1.4\nnot really a pdf\n')
        text = str(caught.exception)
        for leak in ('Traceback', 'Exception', 'pypdf', 'ledgerlens_'):
            self.assertNotIn(leak, text)

    def test_explain_empty_wording(self) -> None:
        self.assertIn('0 bytes', ingestion.explain_empty('a.txt', 0))
        self.assertIn('OCR', ingestion.explain_empty('a.pdf', 100))
        self.assertIn('no data rows', ingestion.explain_empty('a.csv', 100))
        self.assertIn('no data rows', ingestion.explain_empty('a.xlsx', 100))
        self.assertIn('blank', ingestion.explain_empty('a.txt', 100))


class FileNameTests(unittest.TestCase):
    def test_path_parts_are_dropped(self) -> None:
        self.assertEqual(ingestion.clean_filename('../../etc/passwd'), 'passwd')
        self.assertEqual(ingestion.clean_filename('C:\\temp\\..\\evil.txt'), 'evil.txt')

    def test_leading_dots_are_removed(self) -> None:
        self.assertEqual(ingestion.clean_filename('.hidden.txt'), 'hidden.txt')

    def test_odd_characters_become_underscores(self) -> None:
        self.assertEqual(ingestion.clean_filename('a<b>|c?.csv'), 'a_b__c_.csv')

    def test_empty_name_gets_a_fallback(self) -> None:
        self.assertEqual(ingestion.clean_filename(''), 'document')
        self.assertEqual(ingestion.clean_filename('...'), 'document')

    def test_long_names_are_cut(self) -> None:
        self.assertEqual(len(ingestion.clean_filename('a' * 300 + '.txt')), 120)

    def test_the_cleaned_name_is_what_shows_up_as_source(self) -> None:
        documents = read_document('../secret/notes.txt', b'budget review')
        self.assertEqual(documents[0].metadata['source'], 'notes.txt')


class AcceptedFileTests(unittest.TestCase):
    def test_utf8_text(self) -> None:
        documents = read_document('a.txt', 'Meals: \u20ac64 per day'.encode('utf-8'))
        self.assertIn('\u20ac64', documents[0].page_content)

    def test_utf8_with_bom(self) -> None:
        documents = read_document('a.txt', b'\xef\xbb\xbfhello ledger')
        self.assertEqual(documents[0].page_content.strip(), 'hello ledger')

    def test_utf16_text(self) -> None:
        documents = read_document('a.txt', 'quarterly review'.encode('utf-16'))
        self.assertIn('quarterly review', documents[0].page_content)

    def test_windows_1252_text(self) -> None:
        documents = read_document('a.txt', 'caf\u00e9 receipt'.encode('cp1252'))
        self.assertIn('caf\u00e9', documents[0].page_content)

    def test_csv_rows_become_passages_with_sheet_style_row_numbers(self) -> None:
        documents = read_document('claims.csv', CLAIMS_CSV)
        self.assertEqual(len(documents), 3)
        self.assertEqual([d.metadata['row'] for d in documents], [2, 3, 4])
        self.assertIn('department: Finance', documents[0].page_content)
        self.assertEqual(documents[0].metadata['format'], 'csv')
        self.assertEqual(documents[0].metadata['source'], 'claims.csv')

    def test_csv_metadata_is_whitelisted(self) -> None:
        keys = set(read_document('claims.csv', CLAIMS_CSV)[0].metadata)
        self.assertEqual(keys, {'source', 'file_id', 'format', 'row'})

    def test_ragged_csv_still_loads(self) -> None:
        data = b'a,b,c\n1,2\n3,4,5,6\n'
        self.assertEqual(len(read_document('ragged.csv', data)), 2)

    def test_xlsx_rows_carry_sheet_and_spreadsheet_row_number(self) -> None:
        data = workbook_bytes({'Sales': [['Agent', 'Sales'], ['Geeta', 32], ['Dom', 62]]})
        documents = read_document('s.xlsx', data)
        self.assertEqual([d.metadata['row'] for d in documents], [2, 3])
        self.assertEqual(documents[0].metadata['sheet'], 'Sales')
        self.assertEqual(documents[0].page_content, 'Agent: Geeta\nSales: 32')

    def test_xlsx_skips_blank_cells_and_blank_rows(self) -> None:
        data = workbook_bytes({'S': [['a', 'b'], ['x', None], [None, None], ['y', 'z']]})
        documents = read_document('s.xlsx', data)
        self.assertEqual(len(documents), 2)
        self.assertEqual(documents[0].page_content, 'a: x')

    def test_xlsx_reads_every_sheet(self) -> None:
        data = workbook_bytes({'One': [['k'], ['1']], 'Two': [['k'], ['2']]})
        sheets = {d.metadata['sheet'] for d in read_document('m.xlsx', data)}
        self.assertEqual(sheets, {'One', 'Two'})

    def test_upper_case_extension_is_accepted(self) -> None:
        self.assertEqual(len(read_document('CLAIMS.CSV', CLAIMS_CSV)), 3)

    def test_file_id_is_a_stable_content_hash(self) -> None:
        first = read_document('one.csv', CLAIMS_CSV)[0].metadata['file_id']
        second = read_document('other-name.csv', CLAIMS_CSV)[0].metadata['file_id']
        self.assertEqual(first, second)
        self.assertEqual(len(first), 64)


class ChunkingTests(unittest.TestCase):
    def sample(self) -> list[Document]:
        return read_document('memo.txt', ('Claims are due in 45 days. ' * 80).encode('utf-8'))

    def test_chunks_have_ids_and_start_index(self) -> None:
        chunks = ingestion.split_into_chunks(self.sample())
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertEqual(len(chunk.metadata['chunk_id']), 64)
            self.assertIn('start_index', chunk.metadata)
            self.assertLessEqual(len(chunk.page_content), 700)

    def test_chunk_ids_are_stable_between_runs(self) -> None:
        first = [c.metadata['chunk_id'] for c in ingestion.split_into_chunks(self.sample())]
        second = [c.metadata['chunk_id'] for c in ingestion.split_into_chunks(self.sample())]
        self.assertEqual(first, second)

    def test_chunk_ids_are_unique_inside_a_file(self) -> None:
        ids = [c.metadata['chunk_id'] for c in ingestion.split_into_chunks(self.sample())]
        self.assertEqual(len(ids), len(set(ids)))

    def test_changed_content_changes_the_id(self) -> None:
        a = ingestion.split_into_chunks(read_document('a.txt', b'alpha budget'))[0]
        b = ingestion.split_into_chunks(read_document('a.txt', b'omega budget'))[0]
        self.assertNotEqual(a.metadata['chunk_id'], b.metadata['chunk_id'])

    def test_invalid_sizes_are_rejected(self) -> None:
        with self.assertRaises(ValueError):
            ingestion.split_into_chunks(self.sample(), chunk_size=0)
        with self.assertRaises(ValueError):
            ingestion.split_into_chunks(self.sample(), chunk_size=100, chunk_overlap=100)
        with self.assertRaises(ValueError):
            ingestion.split_into_chunks(self.sample(), chunk_size=100, chunk_overlap=-1)

    def test_a_short_table_row_stays_one_passage(self) -> None:
        chunks = ingestion.split_into_chunks(read_document('claims.csv', CLAIMS_CSV))
        self.assertEqual(len(chunks), 3)


class ShippedSampleTests(unittest.TestCase):
    '''The real files in sample_documents/ must load exactly as the docs claim.'''

    @classmethod
    def setUpClass(cls) -> None:
        cls.memo = read_document('Q3_Audit_Memo.txt', (SAMPLES / 'Q3_Audit_Memo.txt').read_bytes())
        cls.claims = read_document('Travel_Expense_Claims.csv', (SAMPLES / 'Travel_Expense_Claims.csv').read_bytes())
        cls.sales = read_document('Q4_Financial_Budget.xlsx', (SAMPLES / 'Q4_Financial_Budget.xlsx').read_bytes())
        cls.policy = read_document('Corporate_Expense_Policy.pdf', (SAMPLES / 'Corporate_Expense_Policy.pdf').read_bytes())

    def test_memo_is_one_passage_with_the_45_day_rule(self) -> None:
        self.assertEqual(len(self.memo), 1)
        self.assertIn('45 calendar days', ' '.join(self.memo[0].page_content.split()))

    def test_claims_csv_has_642_rows(self) -> None:
        self.assertEqual(len(self.claims), 642)
        self.assertEqual(self.claims[0].metadata['row'], 2)
        self.assertIn('short_description: Trade range fly it move.', self.claims[0].page_content)
        self.assertIn('amount: 41825', self.claims[0].page_content)

    def test_sales_sheet_has_9999_rows(self) -> None:
        self.assertEqual(len(self.sales), 9999)
        first = self.sales[0]
        self.assertEqual(first.metadata['sheet'], 'SalesRegion')
        self.assertEqual(first.metadata['row'], 2)
        self.assertIn('Agent: Geeta', first.page_content)
        self.assertIn('State Name: Jharkhand', first.page_content)

    def test_policy_pages_are_numbered_from_one(self) -> None:
        pages = [d.metadata['page'] for d in self.policy]
        self.assertEqual(min(pages), 1)
        self.assertLessEqual(max(pages), 20)

    def test_policy_contains_the_meal_limit(self) -> None:
        text = ' '.join(' '.join(d.page_content.split()) for d in self.policy)
        self.assertIn('$64', text)

    def test_policy_form_pages_only_yield_their_headings(self) -> None:
        # Pages 17-19 are pictures of forms: the parser sees nothing but the page heading (no OCR).
        by_page = {d.metadata['page']: d.page_content for d in self.policy}
        for page in (17, 18, 19):
            with self.subTest(page=page):
                self.assertIn('Appendix', by_page[page])
                self.assertLess(len(by_page[page]), 400)

    def test_the_bad_file_folder_has_all_seven_files(self) -> None:
        names = sorted(p.name for p in BAD_FILES.iterdir())
        self.assertEqual(names, ['accented_names_windows1252.csv', 'corrupted.pdf', 'empty.txt',
                                 'fake_report.docx', 'not_really_excel.xlsx', 'old_excel.xls',
                                 'sample_archive.zip'])

    def test_bad_files_behave_as_documented(self) -> None:
        for name, fragment in (('fake_report.docx', 'Save As'), ('sample_archive.zip', 'Unpack'),
                               ('old_excel.xls', '.xlsx'), ('corrupted.pdf', 'PDF could not be opened'),
                               ('not_really_excel.xlsx', 'Excel file could not be opened')):
            with self.subTest(file=name):
                with self.assertRaises(DocumentRejectedError) as caught:
                    read_document(name, (BAD_FILES / name).read_bytes())
                self.assertIn(fragment, str(caught.exception))
        self.assertEqual(read_document('empty.txt', (BAD_FILES / 'empty.txt').read_bytes()), [])

    def test_windows_1252_csv_gives_three_passages_with_accents(self) -> None:
        documents = read_document('accented_names_windows1252.csv',
                                  (BAD_FILES / 'accented_names_windows1252.csv').read_bytes())
        self.assertEqual(len(documents), 3)
        self.assertIn('Z\u00fcrich', ' '.join(d.page_content for d in documents))


class UploadFolderTests(unittest.TestCase):
    SESSION = 'docs_' + 'a1' * 16

    def test_copy_is_saved_under_the_cleaned_name(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = ingestion.store_upload_copy(self.SESSION, '../x/report.csv', b'abc', Path(folder))
            self.assertEqual(path, Path(folder) / self.SESSION / 'report.csv')
            self.assertEqual(path.read_bytes(), b'abc')

    def test_session_copies_can_be_removed(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            ingestion.store_upload_copy(self.SESSION, 'a.txt', b'1', Path(folder))
            ingestion.remove_session_copies(self.SESSION, Path(folder))
            self.assertFalse((Path(folder) / self.SESSION).exists())

    def test_invalid_session_ids_are_refused(self) -> None:
        for bad in ('', '../etc', 'docs_xyz', 'docs_' + 'A' * 32, 'docs_' + 'a' * 31):
            with self.subTest(session=bad):
                with self.assertRaises(ValueError):
                    ingestion.store_upload_copy(bad, 'a.txt', b'1', Path('.'))

    def test_clearing_old_uploads_keeps_the_gitkeep_marker(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / '.gitkeep').write_bytes(b'')
            ingestion.store_upload_copy(self.SESSION, 'a.txt', b'1', root)
            (root / 'stray.txt').write_bytes(b'x')
            ingestion.clear_old_uploads(root)
            self.assertEqual([p.name for p in root.iterdir()], ['.gitkeep'])


class FakeStore:
    '''Stands in for Chroma: semantic_search only calls similarity_search_with_score.'''

    def __init__(self, pairs):
        self.pairs = pairs
        self.queries: list[str] = []

    def similarity_search_with_score(self, query, k=5):
        self.queries.append(query)
        return self.pairs[:k]


class SearchTests(unittest.TestCase):
    NEAR = Document(page_content='near passage', metadata={'chunk_id': 'n'})
    FAR = Document(page_content='far passage', metadata={'chunk_id': 'f'})

    def test_distance_is_converted_to_similarity(self) -> None:
        hits = retriever.semantic_search(FakeStore([(self.NEAR, 0.2)]), 'meal limit')
        self.assertAlmostEqual(hits[0].similarity, 0.8)

    def test_weak_matches_are_dropped(self) -> None:
        store = FakeStore([(self.NEAR, 0.2), (self.FAR, 0.9)])
        hits = retriever.semantic_search(store, 'meal limit', min_similarity=0.30)
        self.assertEqual([h.document.page_content for h in hits], ['near passage'])

    def test_similarity_is_clamped(self) -> None:
        hits = retriever.semantic_search(FakeStore([(self.NEAR, -0.5), (self.FAR, 3.0)]), 'q', min_similarity=-1)
        self.assertEqual([h.similarity for h in hits], [1.0, -1.0])

    def test_empty_question_returns_nothing_and_asks_no_store(self) -> None:
        store = FakeStore([(self.NEAR, 0.1)])
        self.assertEqual(retriever.semantic_search(store, '   '), [])
        self.assertEqual(store.queries, [])

    def test_the_question_is_stripped_before_searching(self) -> None:
        store = FakeStore([(self.NEAR, 0.1)])
        retriever.semantic_search(store, '  meal limit  ')
        self.assertEqual(store.queries, ['meal limit'])

    def test_k_and_threshold_are_validated(self) -> None:
        store = FakeStore([])
        for bad_k in (0, 21):
            with self.assertRaises(ValueError):
                retriever.semantic_search(store, 'q', k=bad_k)
        for bad_threshold in (-1.5, 1.5):
            with self.assertRaises(ValueError):
                retriever.semantic_search(store, 'q', min_similarity=bad_threshold)

    def test_k_limits_the_number_of_hits(self) -> None:
        store = FakeStore([(self.NEAR, 0.1)] * 10)
        self.assertEqual(len(retriever.semantic_search(store, 'q', k=3)), 3)


class StoreHelperTests(unittest.TestCase):
    def test_collection_names_must_match_the_session_pattern(self) -> None:
        for bad in ('', 'docs', 'docs_123', 'docs_' + 'g' * 32, '../docs_' + 'a' * 32, 'docs_' + 'a' * 33):
            with self.subTest(name=bad):
                with self.assertRaises(ValueError):
                    retriever.initialize_vector_store(bad, Path(tempfile.gettempdir()))

    def test_purge_keeps_the_gitkeep_marker(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / '.gitkeep').write_bytes(b'')
            (root / 'chroma.sqlite3').write_bytes(b'data')
            (root / 'segment').mkdir()
            retriever.purge_stale_vector_data(root)
            self.assertEqual([p.name for p in root.iterdir()], ['.gitkeep'])

    def test_purge_of_a_missing_folder_is_harmless(self) -> None:
        retriever.purge_stale_vector_data(Path(tempfile.gettempdir()) / 'ledgerlens_missing_folder_xyz')

    def test_indexing_uses_chunk_ids_and_batches(self) -> None:
        class Recorder:
            def __init__(self) -> None:
                self.calls: list[tuple[int, list[str]]] = []

            def add_documents(self, docs, ids=None):
                self.calls.append((len(docs), list(ids)))

        chunks = [Document(page_content=f'row {i}', metadata={'chunk_id': f'id{i}'}) for i in range(300)]
        recorder = Recorder()
        progress: list[tuple[int, int]] = []
        total = retriever.index_chunks(recorder, chunks, on_progress=lambda done, all_: progress.append((done, all_)))
        self.assertEqual(total, 300)
        self.assertEqual([size for size, _ in recorder.calls], [128, 128, 44])
        self.assertEqual(recorder.calls[0][1][0], 'id0')
        self.assertEqual(progress[-1], (300, 300))


# ============================================================================= agent, LLM layer, guards
MEAL = Document(page_content='Meals are reimbursed up to $64 per day.\nReceipts are required for every meal.',
                metadata={'source': 'policy.txt', 'chunk_id': 'c1', 'format': 'txt'})


def meal_store() -> FakeStore:
    return FakeStore([(MEAL, 0.2)])  # cosine distance 0.2 -> similarity 0.8


def row_doc(number: int, **fields) -> Document:
    return Document(page_content='\n'.join(f'{key}: {value}' for key, value in fields.items()),
                    metadata={'source': 'sales.xlsx', 'sheet': 'SalesRegion', 'row': number,
                              'chunk_id': f'row{number}', 'format': 'xlsx'})


ROWS = [row_doc(2, Agent='Geeta', Sales=32, Target=41), row_doc(3, Agent='Dom', Sales=62, Target=70),
        row_doc(4, Agent='Akbar', Sales=65, Target=73), row_doc(5, Agent='Vijay', Sales=25, Target=24),
        row_doc(6, Agent='Seeta', Sales=54, Target=50)]


def table_store() -> FakeStore:
    return FakeStore([(row, 0.2 + 0.01 * index) for index, row in enumerate(ROWS)])


class ScriptedModel(GenericFakeChatModel):
    '''A chat model that replays fixed messages and accepts tool binding like a real one.'''

    def bind_tools(self, tools, **kwargs):
        return self


def scripted_factory(final_args: dict, query: str = 'meal limit', calls: list | None = None):
    '''Model factory whose model searches once and then hands in the given final reply.'''
    def factory(provider):
        if calls is not None:
            calls.append(provider.name)
        messages = [
            AIMessage(content='', tool_calls=[{'name': 'search_documents', 'args': {'query': query}, 'id': 'call_1'}]),
            AIMessage(content='', tool_calls=[{'name': 'DraftReply', 'args': final_args, 'id': 'call_2'}]),
        ]
        return ScriptedModel(messages=iter(messages))
    return factory


def failing_factory(error: Exception, calls: list | None = None):
    def factory(provider):
        if calls is not None:
            calls.append(provider.name)
        raise error
    return factory


def by_provider(**factories):
    '''Pick a factory per provider name, so a test can make one model fail and another succeed.'''
    def factory(provider):
        return factories[provider.name](provider)
    return factory


def must_not_be_called(provider):
    raise AssertionError('The language model must not be called here.')


GEM = Provider('gemini', 'Gemini', 'gem-model', 'gem-secret-key-1111')
GRQ = Provider('groq', 'Groq', 'groq-model', 'groq-secret-key-2222')

GOOD_ARGS = {'outcome': 'answered', 'claims': [
    {'text': 'Meals are reimbursed up to $64 per day.',
     'evidence': [{'source_id': 1, 'quote': 'reimbursed up to $64 per day'}]}]}


def ledger_with(*documents: Document, similarity: float = 0.8) -> PassageLedger:
    ledger = PassageLedger()
    for document in documents:
        ledger.add(SearchHit(document, similarity))
    return ledger


def claim(text: str, source_id: int, quote: str) -> AnswerClaim:
    return AnswerClaim(text=text, evidence=[EvidenceItem(source_id=source_id, quote=quote)])


class ValidationTests(unittest.TestCase):
    '''The plain-code citation checker: no model involved, so a model cannot talk its way past it.'''

    def validate(self, draft, ledger):
        return LedgerAgent._validate(draft, ledger)

    def test_a_correct_quote_gives_a_cited_answer(self) -> None:
        draft = DraftReply(outcome='answered', claims=[claim('The limit is $64 a day.', 1, 'up to $64 per day')])
        result = self.validate(draft, ledger_with(MEAL))
        self.assertEqual(result.text, 'The limit is $64 a day. [1]')
        self.assertEqual(result.sources[0]['id'], 1)
        self.assertEqual(result.sources[0]['quotes'], ['up to $64 per day'])
        self.assertEqual(result.sources[0]['similarity'], 0.8)
        self.assertIn('Receipts are required', result.sources[0]['excerpt'])

    def test_whitespace_differences_in_a_quote_are_accepted(self) -> None:
        draft = DraftReply(outcome='answered',
                           claims=[claim('Limit.', 1, 'per day.   Receipts\nare   required')])
        self.assertIn('[1]', self.validate(draft, ledger_with(MEAL)).text)

    def test_curly_quotes_in_the_passage_match_plain_quotes_in_the_quote(self) -> None:
        # Real PDFs print "day’s"; models usually write it back as "day's".
        passage = Document(page_content='up to $64 a day and $25 for a half day’s travel.',
                           metadata={'source': 'p.pdf', 'chunk_id': 'q1', 'format': 'pdf', 'page': 12})
        draft = DraftReply(outcome='answered', claims=[claim('Limit.', 1, "$25 for a half day's travel.")])
        self.assertIn('[1]', self.validate(draft, ledger_with(passage)).text)

    def test_long_dashes_non_breaking_spaces_and_ligatures_are_tolerated(self) -> None:
        passage = Document(page_content='Meals – per day: $64 and the oﬃce rule “no exceptions”.',
                           metadata={'source': 'p.pdf', 'chunk_id': 'q2', 'format': 'pdf'})
        for quote in ('Meals - per day: $64', 'the office rule "no exceptions".'):
            with self.subTest(quote=quote):
                draft = DraftReply(outcome='answered', claims=[claim('Fact.', 1, quote)])
                self.assertIn('[1]', self.validate(draft, ledger_with(passage)).text)

    def test_normalising_punctuation_does_not_let_changed_words_through(self) -> None:
        passage = Document(page_content='half day’s travel costs $25.',
                           metadata={'source': 'p.pdf', 'chunk_id': 'q3', 'format': 'pdf'})
        for quote in ("half day's travel costs $26.", "half days travel costs $25.", "full day's travel costs $25."):
            with self.subTest(quote=quote):
                draft = DraftReply(outcome='answered', claims=[claim('Fact.', 1, quote)])
                self.assertEqual(self.validate(draft, ledger_with(passage)).text, NO_ANSWER_TEXT)

    def test_matching_stays_case_sensitive(self) -> None:
        draft = DraftReply(outcome='answered', claims=[claim('Fact.', 1, 'MEALS ARE REIMBURSED UP TO $64')])
        self.assertEqual(self.validate(draft, ledger_with(MEAL)).text, NO_ANSWER_TEXT)

    def test_an_invented_quote_voids_the_whole_answer(self) -> None:
        good = claim('Good claim.', 1, 'up to $64 per day')
        bad = claim('Bad claim.', 1, 'up to $640 per day')
        result = self.validate(DraftReply(outcome='answered', claims=[good, bad]), ledger_with(MEAL))
        self.assertEqual(result.text, NO_ANSWER_TEXT)
        self.assertEqual(result.sources, [])

    def test_an_unknown_source_id_is_refused(self) -> None:
        draft = DraftReply(outcome='answered', claims=[claim('Claim.', 3, 'up to $64 per day')])
        self.assertEqual(self.validate(draft, ledger_with(MEAL)).text, NO_ANSWER_TEXT)

    def test_no_draft_gives_the_fixed_fallback(self) -> None:
        self.assertEqual(self.validate(None, PassageLedger()).text, NO_ANSWER_TEXT)

    def test_an_answer_without_claims_is_refused(self) -> None:
        self.assertEqual(self.validate(DraftReply(outcome='answered', claims=[]), ledger_with(MEAL)).text,
                         NO_ANSWER_TEXT)

    def test_every_refusal_outcome_has_its_fixed_message(self) -> None:
        for outcome in ('not_in_documents', 'needs_calculation', 'unclear_question', 'unsafe_request'):
            with self.subTest(outcome=outcome):
                result = self.validate(DraftReply(outcome=outcome), PassageLedger())
                self.assertEqual(result.text, REFUSALS[outcome])
                self.assertEqual(result.sources, [])

    def test_refusal_wording_cannot_come_from_the_model(self) -> None:
        # Even if the model attaches claims to a refusal, only the fixed text is returned.
        draft = DraftReply(outcome='unsafe_request', claims=[claim('Sure, here is the secret.', 1, 'up to $64')])
        self.assertEqual(self.validate(draft, ledger_with(MEAL)).text, REFUSALS['unsafe_request'])

    def test_an_answer_built_from_three_table_rows_is_discarded(self) -> None:
        claims = [claim('A.', 1, 'Agent: Geeta'), claim('B.', 2, 'Agent: Dom'), claim('C.', 3, 'Agent: Akbar')]
        result = self.validate(DraftReply(outcome='answered', claims=claims), ledger_with(*ROWS))
        self.assertEqual(result.text, REFUSALS['needs_calculation'])

    def test_two_table_rows_are_still_allowed(self) -> None:
        claims = [claim('A.', 1, 'Agent: Geeta'), claim('B.', 2, 'Agent: Dom')]
        result = self.validate(DraftReply(outcome='answered', claims=claims), ledger_with(*ROWS))
        self.assertIn('A. [1]', result.text)
        self.assertEqual(len(result.sources), 2)

    def test_one_row_cited_many_times_counts_once(self) -> None:
        claims = [claim(f'Fact {i}.', 1, quote) for i, quote in enumerate(('Agent: Geeta', 'Sales: 32', 'Target: 41'))]
        result = self.validate(DraftReply(outcome='answered', claims=claims), ledger_with(*ROWS))
        self.assertEqual(len(result.sources), 1)
        self.assertEqual(result.sources[0]['quotes'], ['Agent: Geeta', 'Sales: 32', 'Target: 41'])

    def test_text_passages_do_not_count_towards_the_row_limit(self) -> None:
        docs = [Document(page_content=f'Rule {i}: keep receipts.', metadata={'source': 'p.txt', 'chunk_id': f't{i}', 'format': 'txt'})
                for i in range(4)]
        claims = [claim(f'Rule {i}.', i + 1, f'Rule {i}: keep receipts.') for i in range(4)]
        result = self.validate(DraftReply(outcome='answered', claims=claims), ledger_with(*docs))
        self.assertEqual(len(result.sources), 4)

    def test_quotes_are_deduplicated_per_source(self) -> None:
        claims = [claim('One.', 1, 'up to $64 per day'), claim('Two.', 1, 'up to $64 per day')]
        result = self.validate(DraftReply(outcome='answered', claims=claims), ledger_with(MEAL))
        self.assertEqual(result.sources[0]['quotes'], ['up to $64 per day'])


class PassageLedgerTests(unittest.TestCase):
    def test_passages_are_numbered_from_one(self) -> None:
        ledger = ledger_with(*ROWS[:3])
        self.assertEqual([ledger.get(n).metadata['row'] for n in (1, 2, 3)], [2, 3, 4])
        self.assertIsNone(ledger.get(0))
        self.assertIsNone(ledger.get(4))

    def test_the_same_chunk_keeps_its_number(self) -> None:
        ledger = PassageLedger()
        first = ledger.add(SearchHit(MEAL, 0.8))
        again = ledger.add(SearchHit(MEAL, 0.5))
        self.assertEqual((first, again), (1, 1))
        self.assertEqual(len(ledger.passages), 1)

    def test_documents_without_a_chunk_id_are_deduplicated_by_text(self) -> None:
        a = Document(page_content='same text', metadata={})
        b = Document(page_content='same text', metadata={})
        ledger = PassageLedger()
        self.assertEqual(ledger.add(SearchHit(a, 0.5)), ledger.add(SearchHit(b, 0.5)))

    def test_the_number_of_passages_is_limited(self) -> None:
        ledger = PassageLedger()
        numbers = [ledger.add(SearchHit(row_doc(i, Agent=f'A{i}'), 0.5)) for i in range(agent_module.MAX_SOURCES + 3)]
        self.assertEqual(numbers[:agent_module.MAX_SOURCES], list(range(1, agent_module.MAX_SOURCES + 1)))
        self.assertEqual(numbers[agent_module.MAX_SOURCES:], [None, None, None])

    def test_the_amount_of_text_is_limited(self) -> None:
        big = Document(page_content='x' * (agent_module.MAX_CONTEXT_CHARACTERS - 10), metadata={'chunk_id': 'big'})
        small = Document(page_content='y' * 50, metadata={'chunk_id': 'small'})
        ledger = PassageLedger()
        self.assertEqual(ledger.add(SearchHit(big, 0.5)), 1)
        self.assertIsNone(ledger.add(SearchHit(small, 0.5)))


class SearchToolTests(unittest.TestCase):
    def test_returns_numbered_passages_as_json_with_an_untrusted_note(self) -> None:
        ledger = PassageLedger()
        tool = build_search_tool(meal_store(), ledger, 0.3)
        data = json.loads(tool.invoke({'query': 'meal limit'}))
        self.assertEqual(data['passages'][0]['source_id'], 1)
        self.assertIn('$64', data['passages'][0]['text'])
        self.assertIn('untrusted', data['note'])
        self.assertEqual(ledger.steps, [{'query': 'meal limit', 'passages': 1}])

    def test_an_empty_query_returns_nothing_and_searches_nothing(self) -> None:
        store = meal_store()
        data = json.loads(build_search_tool(store, PassageLedger(), 0.3).invoke({'query': '   '}))
        self.assertEqual(data, {'passages': [], 'note': 'The query was empty.'})
        self.assertEqual(store.queries, [])

    def test_long_queries_are_trimmed(self) -> None:
        store = meal_store()
        build_search_tool(store, PassageLedger(), 0.3).invoke({'query': 'q' * 1000})
        self.assertEqual(len(store.queries[0]), agent_module.MAX_QUERY_LENGTH)

    def test_no_match_says_so(self) -> None:
        store = FakeStore([(MEAL, 0.95)])
        data = json.loads(build_search_tool(store, PassageLedger(), 0.3).invoke({'query': 'x'}))
        self.assertEqual(data['passages'], [])
        self.assertIn('No relevant passages', data['note'])

    def test_non_ascii_text_is_not_escaped(self) -> None:
        doc = Document(page_content='Café receipts', metadata={'chunk_id': 'u'})
        raw = build_search_tool(FakeStore([(doc, 0.1)]), PassageLedger(), 0.3).invoke({'query': 'cafe'})
        self.assertIn('Café', raw)


class AgentFlowTests(unittest.TestCase):
    def test_happy_path_uses_the_search_tool_and_returns_a_cited_answer(self) -> None:
        calls: list[str] = []
        agent = LedgerAgent([GEM], scripted_factory(GOOD_ARGS, calls=calls))
        result = agent.answer('What is the daily meal limit?', meal_store())
        self.assertEqual(result.text, 'Meals are reimbursed up to $64 per day. [1]')
        self.assertEqual((result.provider, result.model, result.fell_back), ('Gemini', 'gem-model', False))
        self.assertEqual(result.steps, [{'query': 'meal limit', 'passages': 1}])
        self.assertEqual(result.note, '')
        self.assertEqual(calls, ['gemini'])
        self.assertEqual(agent.last_provider, 'gemini')

    def test_a_model_that_invents_a_quote_is_caught(self) -> None:
        bad = {'outcome': 'answered', 'claims': [{'text': 'Meals are unlimited.',
                                                  'evidence': [{'source_id': 1, 'quote': 'meals are unlimited for everyone'}]}]}
        result = LedgerAgent([GEM], scripted_factory(bad)).answer('Meal limit?', meal_store())
        self.assertEqual(result.text, NO_ANSWER_TEXT)

    def test_a_model_can_decline_with_the_fixed_message(self) -> None:
        for outcome in ('not_in_documents', 'needs_calculation', 'unclear_question', 'unsafe_request'):
            with self.subTest(outcome=outcome):
                agent = LedgerAgent([GEM], scripted_factory({'outcome': outcome, 'claims': []}))
                self.assertEqual(agent.answer('Meal limit?', meal_store()).text, REFUSALS[outcome])

    def test_no_relevant_passage_never_calls_the_model(self) -> None:
        agent = LedgerAgent([GEM], must_not_be_called)
        result = agent.answer('Who won the world cup?', FakeStore([(MEAL, 0.95)]))
        self.assertEqual(result.text, NO_ANSWER_TEXT)
        self.assertIn('no language model was called', result.note)
        self.assertEqual(result.provider, '')

    def test_empty_question_asks_for_input(self) -> None:
        result = LedgerAgent([GEM], must_not_be_called).answer('   ', meal_store())
        self.assertIn('type a question', result.text)

    def test_overlong_question_asks_for_a_shorter_one(self) -> None:
        result = LedgerAgent([GEM], must_not_be_called).answer('x' * 2001, meal_store())
        self.assertIn('2000', result.text)

    def test_the_question_is_stripped(self) -> None:
        store = meal_store()
        LedgerAgent([GEM], scripted_factory(GOOD_ARGS)).answer('   Meal limit?  ', store)
        self.assertEqual(store.queries[0], 'Meal limit?')

    def test_at_least_one_provider_is_required(self) -> None:
        with self.assertRaises(ValueError):
            LedgerAgent([])

    def test_falls_back_from_the_first_model_to_the_second(self) -> None:
        factory = by_provider(gemini=failing_factory(RuntimeError('503 UNAVAILABLE high demand')),
                              groq=scripted_factory(GOOD_ARGS))
        result = LedgerAgent([GEM, GRQ], factory).answer('Meal limit?', meal_store())
        self.assertEqual((result.provider, result.fell_back), ('Groq', True))
        self.assertIn('$64', result.text)

    def test_a_failed_model_is_tried_last_next_time(self) -> None:
        calls: list[str] = []
        factory = by_provider(gemini=failing_factory(RuntimeError('503 UNAVAILABLE'), calls),
                              groq=scripted_factory(GOOD_ARGS, calls=calls))
        agent = LedgerAgent([GEM, GRQ], factory)
        agent.answer('First?', meal_store())
        agent.answer('Second?', meal_store())
        self.assertEqual(calls, ['gemini', 'groq', 'groq'])  # the second question went straight to Groq

    def test_a_failed_model_is_not_skipped_when_it_is_the_only_one_left(self) -> None:
        calls: list[str] = []
        agent = LedgerAgent([GEM, GRQ], by_provider(gemini=failing_factory(RuntimeError('503'), calls),
                                                    groq=failing_factory(RuntimeError('503'), calls)))
        for _ in range(2):
            with self.assertRaises(AgentFailure):
                agent.answer('Meal limit?', meal_store())
        self.assertEqual(len(calls), 4)

    def test_a_working_model_clears_its_block(self) -> None:
        state = {'fail': True}

        def flaky(provider):
            if state['fail']:
                raise RuntimeError('503 UNAVAILABLE')
            return scripted_factory(GOOD_ARGS)(provider)

        agent = LedgerAgent([GEM], flaky)
        with self.assertRaises(AgentFailure):
            agent.answer('Meal limit?', meal_store())
        self.assertIn('gemini', agent._blocked_until)
        state['fail'] = False
        agent.answer('Meal limit?', meal_store())
        self.assertNotIn('gemini', agent._blocked_until)

    def test_all_models_failing_raises_a_safe_error(self) -> None:
        factory = by_provider(gemini=failing_factory(RuntimeError(f'boom {GEM.api_key} kaboom-trace')),
                              groq=failing_factory(ValueError('401 invalid api key abc')))
        with self.assertRaises(AgentFailure) as caught:
            LedgerAgent([GEM, GRQ], factory).answer('Meal limit?', meal_store())
        text = str(caught.exception)
        self.assertTrue(text.startswith('No language model could answer right now.'))
        self.assertIn('Gemini:', text)
        self.assertIn('Groq: the API key was not accepted', text)
        for leak in ('boom', 'kaboom-trace', GEM.api_key, 'abc'):
            self.assertNotIn(leak, text)

    def test_a_search_failure_becomes_a_friendly_error(self) -> None:
        class BrokenStore:
            def similarity_search_with_score(self, query, k=5):
                raise RuntimeError('chroma exploded with /private/path')

        with self.assertRaises(AgentFailure) as caught:
            LedgerAgent([GEM], must_not_be_called).answer('Meal limit?', BrokenStore())
        self.assertIn('Clear the knowledge base', str(caught.exception))
        self.assertNotIn('exploded', str(caught.exception))

    def test_a_model_that_never_returns_a_draft_gives_the_fallback(self) -> None:
        def chatty(provider):
            return ScriptedModel(messages=iter([AIMessage(content='I think the limit is 64.')]))

        result = LedgerAgent([GEM], chatty).answer('Meal limit?', meal_store())
        self.assertEqual(result.text, NO_ANSWER_TEXT)


class LlmHelperTests(unittest.TestCase):
    ENV = dict(GOOGLE_API_KEY='AIzaRealLookingGeminiKey', GROQ_API_KEY='gsk_RealLookingGroqKey')

    def test_real_keys_are_told_apart_from_placeholders(self) -> None:
        for good in ('AIzaSyRealKey123', 'gsk_abcdefgh'):
            self.assertTrue(llm.looks_like_real_key(good), good)
        for bad in (None, '', '   ', 'your_gemini_api_key_here', 'paste_key_here', 'CHANGEME', 'xxxx1234', 'Placeholder'):
            self.assertFalse(llm.looks_like_real_key(bad), bad)

    def test_default_model_names(self) -> None:
        with mock.patch.dict(os.environ, {}, clear=True):
            self.assertEqual(llm.gemini_model(), 'gemini-3.6-flash')
            self.assertEqual(llm.gemini_spare_model(), 'gemini-3.5-flash-lite')
            self.assertEqual(llm.groq_model(), 'qwen/qwen3.8-27b')
            self.assertEqual(llm.groq_spare_model(), 'openai/gpt-oss-20b')

    def test_environment_overrides_are_used(self) -> None:
        env = {'GEMINI_MODEL': 'my-gemini', 'GROQ_MODEL': 'my/groq', 'GEMINI_BACKUP_MODEL': 'my-gemini-2',
               'GROQ_BACKUP_MODEL': 'my/groq-2'}
        with mock.patch.dict(os.environ, env, clear=True):
            self.assertEqual(llm.gemini_model(), 'my-gemini')
            self.assertEqual(llm.groq_model(), 'my/groq')
            self.assertEqual(llm.gemini_spare_model(), 'my-gemini-2')
            self.assertEqual(llm.groq_spare_model(), 'my/groq-2')

    def test_an_explicit_name_beats_the_environment(self) -> None:
        with mock.patch.dict(os.environ, {'GEMINI_MODEL': 'from-env'}, clear=True):
            self.assertEqual(llm.gemini_model('explicit'), 'explicit')

    def test_blank_and_retired_names_fall_back_to_the_default(self) -> None:
        with mock.patch.dict(os.environ, {'GEMINI_MODEL': '   ', 'GROQ_MODEL': 'llama-3.3-70b-versatile',
                                          'GROQ_BACKUP_MODEL': 'openai/gpt-oss-120b',
                                          'GEMINI_BACKUP_MODEL': 'gemini-2.5-flash'}, clear=True):
            self.assertEqual(llm.gemini_model(), llm.GEMINI_MAIN)
            self.assertEqual(llm.groq_model(), llm.GROQ_MAIN)
            self.assertEqual(llm.groq_spare_model(), llm.GROQ_SPARE)
            self.assertEqual(llm.gemini_spare_model(), llm.GEMINI_SPARE)

    def test_no_keys_means_no_providers(self) -> None:
        with mock.patch.dict(os.environ, {}, clear=True):
            self.assertEqual(llm.active_providers(), [])

    def test_placeholder_keys_are_ignored(self) -> None:
        env = {'GOOGLE_API_KEY': 'your_gemini_api_key_here', 'GROQ_API_KEY': 'your_groq_api_key_here'}
        with mock.patch.dict(os.environ, env, clear=True):
            self.assertEqual(llm.active_providers(), [])

    def test_full_chain_order(self) -> None:
        with mock.patch.dict(os.environ, self.ENV, clear=True):
            names = [p.name for p in llm.active_providers()]
        self.assertEqual(names, ['gemini', 'groq', 'groq-spare', 'gemini-spare'])

    def test_one_company_gives_two_models_on_the_same_key(self) -> None:
        with mock.patch.dict(os.environ, {'GOOGLE_API_KEY': self.ENV['GOOGLE_API_KEY']}, clear=True):
            chain = llm.active_providers()
        self.assertEqual([p.name for p in chain], ['gemini', 'gemini-spare'])
        self.assertEqual({p.api_key for p in chain}, {self.ENV['GOOGLE_API_KEY']})
        self.assertNotEqual(chain[0].model, chain[1].model)

    def test_a_spare_with_the_same_model_name_is_left_out(self) -> None:
        env = dict(self.ENV, GEMINI_BACKUP_MODEL='gemini-3.6-flash', GROQ_BACKUP_MODEL='qwen/qwen3.8-27b')
        with mock.patch.dict(os.environ, env, clear=True):
            self.assertEqual([p.name for p in llm.active_providers()], ['gemini', 'groq'])

    def test_primary_llm_groq_swaps_the_first_two(self) -> None:
        for value in ('groq', 'GROQ', ' Groq '):
            with self.subTest(value=value):
                with mock.patch.dict(os.environ, dict(self.ENV, PRIMARY_LLM=value), clear=True):
                    names = [p.name for p in llm.active_providers()]
                self.assertEqual(names, ['groq', 'gemini', 'groq-spare', 'gemini-spare'])

    def test_a_sidebar_key_beats_the_environment_key(self) -> None:
        with mock.patch.dict(os.environ, self.ENV, clear=True):
            chain = llm.active_providers(gemini_key='sidebar-gemini-key-999')
        self.assertEqual(chain[0].api_key, 'sidebar-gemini-key-999')
        self.assertEqual(chain[1].api_key, self.ENV['GROQ_API_KEY'])

    def test_a_placeholder_in_the_sidebar_does_not_hide_the_environment_key(self) -> None:
        with mock.patch.dict(os.environ, self.ENV, clear=True):
            chain = llm.active_providers(gemini_key='your_key')
        self.assertEqual(chain[0].api_key, self.ENV['GOOGLE_API_KEY'])

    def test_repr_never_shows_the_key(self) -> None:
        self.assertNotIn(GEM.api_key, repr(GEM))
        self.assertNotIn('secret', repr(GEM))

    def test_unknown_provider_cannot_be_built(self) -> None:
        with self.assertRaises(ValueError):
            llm.build_chat_model(Provider('mystery', 'Mystery', 'm', 'k'))

    def test_error_descriptions(self) -> None:
        cases = [
            (RuntimeError('HTTP 401 Unauthorized'), 'API key was not accepted'),
            (RuntimeError('API key not valid. Please pass a valid API key.'), 'API key was not accepted'),
            (RuntimeError('429 RESOURCE_EXHAUSTED'), 'free usage limit'),
            (RuntimeError('Rate limit reached for model'), 'free usage limit'),
            (RuntimeError('404 model_not_found'), 'model name was not recognised'),
            (RuntimeError('tool_use_failed: Failed to call a function'), 'unusable format'),
            (RuntimeError('503 UNAVAILABLE: high demand'), 'busy or cannot be reached'),
            (TimeoutError('request timed out'), 'busy or cannot be reached'),
            (RuntimeError('something nobody expected'), 'unexpected error'),
        ]
        for exc, fragment in cases:
            with self.subTest(error=str(exc)):
                self.assertIn(fragment, llm.describe_error(exc))

    def test_descriptions_never_echo_the_raw_message(self) -> None:
        self.assertNotIn('hunter2', llm.describe_error(RuntimeError('401 bad key hunter2')))
        self.assertNotIn('hunter2', llm.describe_error(RuntimeError('weird failure hunter2')))

    def test_redact_hides_long_secrets_only(self) -> None:
        self.assertEqual(llm.redact('key=abcdefghij ok', ['abcdefghij']), 'key=*** ok')
        self.assertEqual(llm.redact('pin=1234', ['1234']), 'pin=1234')
        self.assertEqual(llm.redact('a-secret-one and b-secret-two', ['a-secret-one', 'b-secret-two']), '*** and ***')
        self.assertEqual(llm.redact('nothing here', ['', 'abcdefghij']), 'nothing here')

    def test_status_check_reports_a_rejected_gemini_key(self) -> None:
        with mock.patch('google.genai.Client') as client:
            client.return_value.models.get.side_effect = RuntimeError('400 API key not valid')
            works, reason = llm.check_provider(GEM)
        self.assertFalse(works)
        self.assertIn('API key was not accepted', reason)

    def test_status_check_accepts_a_working_gemini_key_without_generating(self) -> None:
        with mock.patch('google.genai.Client') as client:
            works, reason = llm.check_provider(GEM)
        self.assertEqual((works, reason), (True, 'key works'))
        client.return_value.models.get.assert_called_once()
        client.return_value.models.generate_content.assert_not_called()

    def test_status_check_notices_a_missing_groq_model(self) -> None:
        listing = mock.Mock(data=[mock.Mock(id='some-other-model')])
        with mock.patch('groq.Groq') as groq:
            groq.return_value.models.list.return_value = listing
            works, reason = llm.check_provider(GRQ)
        self.assertEqual((works, reason), (False, 'the model name was not recognised for this key'))

    def test_status_check_accepts_a_listed_groq_model(self) -> None:
        listing = mock.Mock(data=[mock.Mock(id='groq-model')])
        with mock.patch('groq.Groq') as groq:
            groq.return_value.models.list.return_value = listing
            self.assertEqual(llm.check_provider(GRQ), (True, 'key works'))


class LocalGuardTests(unittest.TestCase):
    '''Plain-code refusals that must happen BEFORE any language model is called.'''

    AGGREGATIONS = [
        'What is the total amount of all expense requests in the claims file?',
        'How many expense requests were declined?',
        'Which region has the highest total sales?',
        'What is the average sales per agent?',
        'Show the top 5 departments by amount.',
        'Give me the count of requests per category.',
        'Which agent has the lowest sales?',
        'Add up all the sales in Goa.',
        'What are the total sales by region?',
        'Give me a breakdown of requests by state.',
    ]
    LOOKUPS = [
        'What was the Sales figure for agent Geeta in Jharkhand on 2015-07-20?',
        'What is the amount of the request described as "Trade range fly it move."?',
        'What is the state of the request "Area personal despite."?',
        'Which manager is listed for agent Dom in Andhra Pradesh on 2015-12-08?',
        'What is the Target for agent Vijay?',
        'What is the processing time in hours of the request "Day none teacher face mention race sell."?',
        'What is the invoice number of the Acme order?',
        'Which department and category does request R-1 belong to?',
    ]
    ROW_ARGS = {'outcome': 'answered', 'claims': [
        {'text': 'Geeta sold 32.', 'evidence': [{'source_id': 1, 'quote': 'Agent: Geeta'}]}]}

    def test_aggregation_questions_are_refused_with_zero_model_calls(self) -> None:
        agent = LedgerAgent([GEM, GRQ], must_not_be_called)
        for question in self.AGGREGATIONS:
            with self.subTest(question=question):
                result = agent.answer(question, table_store())
                self.assertEqual(result.text, REFUSALS['needs_calculation'])
                self.assertIn('no language model was called', result.note)
                self.assertEqual(result.provider, '')

    def test_look_alike_lookups_still_reach_the_model(self) -> None:
        for question in self.LOOKUPS:
            with self.subTest(question=question):
                calls: list[str] = []
                agent = LedgerAgent([GEM], scripted_factory(self.ROW_ARGS, query='agent', calls=calls))
                result = agent.answer(question, table_store())
                self.assertEqual(calls, ['gemini'])
                self.assertEqual(result.note, '')
                self.assertIn('Geeta sold 32.', result.text)

    def test_the_same_words_are_fine_when_the_evidence_is_not_a_table(self) -> None:
        for question in ('How many working days does the policy give for a cash advance report?',
                         'What is the maximum meal reimbursement?', 'Compare the receipt rules.'):
            with self.subTest(question=question):
                calls: list[str] = []
                LedgerAgent([GEM], scripted_factory(GOOD_ARGS, calls=calls)).answer(question, meal_store())
                self.assertEqual(calls, ['gemini'])

    def test_a_mostly_table_result_needs_more_than_half_table_rows(self) -> None:
        mixed = FakeStore([(MEAL, 0.2), (MEAL.model_copy(update={'metadata': dict(MEAL.metadata, chunk_id='c2')}), 0.21),
                           (ROWS[0], 0.22), (ROWS[1], 0.23)])
        calls: list[str] = []
        LedgerAgent([GEM], scripted_factory(GOOD_ARGS, calls=calls)).answer('What is the highest limit?', mixed)
        self.assertEqual(calls, ['gemini'])  # 2 of 4 is not "more than half"

    def test_the_cue_ignores_column_names_on_their_own(self) -> None:
        for text in ('What is the Total row?', 'invoice number', 'What is the amount of the request?',
                     'the state of the request', 'the Target column', 'type of expense', 'teacher each'):
            with self.subTest(text=text):
                self.assertIsNone(agent_module.AGGREGATION_CUE.search(text))

    def test_the_cue_finds_aggregation_wording(self) -> None:
        for text in ('how many', 'total amount of all claims', 'totals per department', 'sum of amounts',
                     'average', 'median', 'percentage of declined', 'highest', 'top 10 agents', 'rank the regions',
                     'compare departments', 'list all requests', 'by department', 'number of claims', 'grand total'):
            with self.subTest(text=text):
                self.assertIsNotNone(agent_module.AGGREGATION_CUE.search(text))

    def test_the_no_evidence_refusal_explains_itself(self) -> None:
        result = LedgerAgent([GEM], must_not_be_called).answer('Who won the world cup?', FakeStore([]))
        self.assertEqual(result.text, NO_ANSWER_TEXT)
        self.assertTrue(result.note.startswith('Local check:'))

    def test_an_answer_glued_from_many_rows_is_discarded(self) -> None:
        args = {'outcome': 'answered', 'claims': [
            {'text': 'Geeta.', 'evidence': [{'source_id': 1, 'quote': 'Agent: Geeta'}]},
            {'text': 'Dom.', 'evidence': [{'source_id': 2, 'quote': 'Agent: Dom'}]},
            {'text': 'Akbar.', 'evidence': [{'source_id': 3, 'quote': 'Agent: Akbar'}]}]}
        result = LedgerAgent([GEM], scripted_factory(args, query='agents')).answer('Tell me about three agents.', table_store())
        self.assertEqual(result.text, REFUSALS['needs_calculation'])

    def test_a_malformed_tool_call_is_retried_once_on_the_same_model(self) -> None:
        calls: list[str] = []
        good = scripted_factory(GOOD_ARGS)

        def glitchy_then_fine(provider):
            calls.append(provider.name)
            if len(calls) == 1:
                raise RuntimeError("400 tool_use_failed: Failed to call a function. failed_generation: '<function=x>'")
            return good(provider)

        result = LedgerAgent([GEM, GRQ], glitchy_then_fine).answer('Meal limit?', meal_store())
        self.assertEqual(calls, ['gemini', 'gemini'])
        self.assertEqual((result.provider, result.fell_back), ('Gemini', False))

    def test_two_malformed_tool_calls_move_on_to_the_next_model(self) -> None:
        calls: list[str] = []
        glitch = RuntimeError('tool call validation failed: invalid arguments')
        factory = by_provider(gemini=failing_factory(glitch, calls), groq=scripted_factory(GOOD_ARGS, calls=calls))
        result = LedgerAgent([GEM, GRQ], factory).answer('Meal limit?', meal_store())
        self.assertEqual(calls, ['gemini', 'gemini', 'groq'])
        self.assertEqual((result.provider, result.fell_back), ('Groq', True))


class RateLimitTests(unittest.TestCase):
    def test_retry_hints_are_parsed(self) -> None:
        cases = [
            ('Rate limit reached. Please try again in 1m3.5s.', 63.5),
            ('quota exceeded, retry in 51s', 51.0),
            ('Please try again in 250ms', 0.25),
            ("{'retryDelay': '34s'}", 34.0),
            ('try again in 2h', 7200.0),
        ]
        for text, seconds in cases:
            with self.subTest(text=text):
                self.assertAlmostEqual(llm.retry_delay_seconds(RuntimeError(text)), seconds)

    def test_no_hint_means_none(self) -> None:
        self.assertIsNone(llm.retry_delay_seconds(RuntimeError('429 too many requests')))

    def test_daily_limits_are_recognised(self) -> None:
        for text in ('Limit exceeded: requests per day', 'tokens per day (TPD) reached', 'GenerateRequestsPerDay quota'):
            self.assertTrue(llm.is_daily_limit(RuntimeError(text)), text)
        self.assertFalse(llm.is_daily_limit(RuntimeError('tokens per minute reached')))

    def test_quota_errors_are_recognised(self) -> None:
        for text in ('429', 'RESOURCE_EXHAUSTED', 'You exceeded your current quota', 'Too Many Requests'):
            self.assertTrue(llm.is_quota_error(RuntimeError(text)), text)
        self.assertFalse(llm.is_quota_error(RuntimeError('503 unavailable')))

    def test_malformed_tool_calls_are_recognised(self) -> None:
        for text in ('tool_use_failed', 'Tool call validation failed', 'Failed to call a function', 'failed_generation'):
            self.assertTrue(llm.is_malformed_tool_call(RuntimeError(text)), text)
        self.assertFalse(llm.is_malformed_tool_call(RuntimeError('timeout')))

    def test_numbers_inside_other_numbers_are_not_status_codes(self) -> None:
        self.assertNotEqual(llm.describe_error(RuntimeError('Requested 4013 tokens, limit 6000')),
                            'the API key was not accepted (it may be invalid, expired or not allowed)')
        self.assertFalse(llm.is_quota_error(RuntimeError('score 0.429 and id 14290')))
        self.assertTrue(llm._has_code('error 401 here', 401))
        self.assertFalse(llm._has_code('4013', 401))
        self.assertFalse(llm._has_code('0.401', 401))

    def test_cooldown_policy(self) -> None:
        cooldown = LedgerAgent._cooldown_seconds
        self.assertEqual(cooldown(RuntimeError('503 unavailable')), agent_module.PROVIDER_COOLDOWN_SECONDS)
        self.assertEqual(cooldown(RuntimeError('429 quota')), agent_module.UNKNOWN_LIMIT_COOLDOWN_SECONDS)
        self.assertEqual(cooldown(RuntimeError('429 try again in 20s')), 22)
        self.assertEqual(cooldown(RuntimeError('429 try again in 1s')), 5)
        self.assertEqual(cooldown(RuntimeError('429 try again in 5h')), agent_module.QUOTA_COOLDOWN_SECONDS)
        self.assertEqual(cooldown(RuntimeError('429 requests per day exceeded')), agent_module.QUOTA_COOLDOWN_SECONDS)
        self.assertEqual(cooldown(RuntimeError('429 requests per day, retry in 7200s')), 7200)

    def test_a_quota_failure_blocks_longer_than_an_outage(self) -> None:
        clock = {'now': 1000.0}
        with mock.patch.object(agent_module, '_now', lambda: clock['now']):
            outage = LedgerAgent([GEM], failing_factory(RuntimeError('503 unavailable')))
            quota = LedgerAgent([GEM], failing_factory(RuntimeError('429 requests per day')))
            for agent in (outage, quota):
                with self.assertRaises(AgentFailure):
                    agent.answer('Meal limit?', meal_store())
            self.assertGreater(quota._blocked_until['gemini'], outage._blocked_until['gemini'])
            self.assertEqual(outage._blocked_until['gemini'], 1000.0 + agent_module.PROVIDER_COOLDOWN_SECONDS)

    def test_a_block_expires(self) -> None:
        clock = {'now': 0.0}
        with mock.patch.object(agent_module, '_now', lambda: clock['now']):
            agent = LedgerAgent([GEM, GRQ], failing_factory(RuntimeError('503')))
            agent._blocked_until['gemini'] = 50.0
            self.assertEqual([p.name for p in agent._ordered_providers()], ['groq', 'gemini'])
            clock['now'] = 51.0
            self.assertEqual([p.name for p in agent._ordered_providers()], ['gemini', 'groq'])


if __name__ == '__main__':
    unittest.main()
